#!/usr/bin/env python3
"""Neue Präsentation mit R+V-Startseite aus der Vorlage.
Aufruf: python3 scripts/new_deck.py --title "Titel" --subtitle "Untertitel" --line "Wiesbaden, 02.10.2026, Anna Beispiel" \
                                    --image bild.jpg --out deck.pptx
- Startseite: Zeile (weiß), Haupttitel (orange), Untertitel (weiß), Bild rechts (füllt die Fläche, wird zugeschnitten).
- Der Titel steht außerdem in der Fußzeile aller Inhaltsseiten (unten links, mit Foliennummer und Datum).
Das Bild muss zum Thema der Präsentation passen: geliefertes Bild verwenden; ohne Bild bleibt der Platzhalter leer (kein Bild erstellen, siehe docs/praesentationen.md)."""
import argparse, pathlib, subprocess, sys, tempfile
from pptx import Presentation

ap = argparse.ArgumentParser()
ap.add_argument("--title", required=True)
ap.add_argument("--subtitle", default="")
ap.add_argument("--line", default="", help="Ort, Datum, Referentin oder Referent")
ap.add_argument("--image", default=None, help="Bild zum Thema der Präsentation; ohne Angabe bleibt der Bildplatzhalter leer")
ap.add_argument("--out", required=True)
a = ap.parse_args()

here = pathlib.Path(__file__).resolve().parent
with tempfile.TemporaryDirectory() as td:
    base = pathlib.Path(td) / "base.pptx"
    subprocess.run([sys.executable, str(here / "build_pptx.py"), "--title", a.title, "--out", str(base)], check=True, stdout=subprocess.DEVNULL)
    prs = Presentation(base)
    s = prs.slides.add_slide(prs.slide_layouts.get_by_name("Startseite"))
    for ph in s.placeholders:
        t = ph.placeholder_format.type
        i = ph.placeholder_format.idx
        if i == 0: ph.text = a.title
        elif i == 1: ph.text = a.subtitle
        elif i == 13: ph.text = a.line
        elif i == 14:
            if a.image: ph.insert_picture(a.image)
            else: print("Hinweis: kein Bild geliefert, Bildplatzhalter bleibt leer", file=sys.stderr)
    prs.save(a.out)
print(a.out)
