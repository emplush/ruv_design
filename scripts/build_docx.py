#!/usr/bin/env python3
"""Erzeugt templates/ruv.docx (A4) aus tokens/tokens.json. Schrift: Arial (Office-Fallback laut Markenportal)."""
import json, pathlib
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root = pathlib.Path(__file__).resolve().parent.parent
T = json.loads((root / "tokens/tokens.json").read_text())
C = {k: v["value"].lstrip("#").upper() for k, v in T["color"].items()}
FONT = T["font"]["office-fallback"]["value"]
rgb = lambda k: RGBColor.from_string(C[k])

d = Document()
sec = d.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.5)
sec.top_margin, sec.bottom_margin = Cm(3), Cm(2.5)

def style(name, size, bold=False, color="primary", before=0, after=6, italic=False):
    s = d.styles[name]
    s.font.name = FONT; s.font.size = Pt(size); s.font.bold = bold; s.font.italic = italic
    s.font.color.rgb = rgb(color)
    rf = s.element.get_or_add_rPr().get_or_add_rFonts()
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rf.set(qn(a), FONT)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:cstheme", "w:eastAsiaTheme"):
        rf.attrib.pop(qn(a), None)
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    return s

n = style("Normal", 11); n.paragraph_format.line_spacing = 1.25
style("Title", 28, True, after=12)
style("Subtitle", 14, False, "orange-dark", after=12)
style("Heading 1", 20, True, before=18, after=6)
style("Heading 2", 15, True, before=14, after=4)
style("Heading 3", 12, True, "mint-dark", before=10, after=3)
style("List Bullet", 11); style("List Number", 11)
style("Quote", 12, False, "primary", italic=True)
# Titel: orange Linie darunter
pPr = d.styles["Title"].element.get_or_add_pPr()
b = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
for k, v in (("w:val", "single"), ("w:sz", "12"), ("w:space", "6"), ("w:color", C["orange-dark"])):
    bt.set(qn(k), v)
b.append(bt); pPr.append(b)

# Kopf- / Fusszeile
hp = sec.header.paragraphs[0]
hp.add_run().add_picture(str(root / "assets/logo/ruv-logo_ohne-claim_positiv.png"), height=Cm(1.6))
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = fp.add_run("Seite "); r.font.size = Pt(9); r.font.color.rgb = rgb("grey-5")
def field(p, code):
    run = p.add_run(); run.font.size = Pt(9); run.font.color.rgb = rgb("grey-5")
    for t, txt in (("begin", None), (None, code), ("end", None)):
        if t:
            e = OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"), t)
        else:
            e = OxmlElement("w:instrText"); e.set(qn("xml:space"), "preserve"); e.text = txt
        run._r.append(e)
field(fp, "PAGE")

# Beispielinhalt
d.add_paragraph("Dokumenttitel", style="Title")
d.add_paragraph("Topline oder Untertitel", style="Subtitle")
d.add_paragraph("Vielen Dank für Ihre Nachricht. Wir kümmern uns persönlich darum und melden uns in den nächsten Tagen bei Ihnen. Hervorhebungen setzen wir fett.")
d.add_heading("Überschrift 1", 1)
d.add_paragraph("Gerne beraten wir Sie auch telefonisch.")
d.add_heading("Überschrift 2", 2)
d.add_paragraph("Schicken Sie uns bitte noch Ihre Unterlagen.", style="List Bullet"); d.add_paragraph("Wir prüfen sie sorgfältig und antworten Ihnen schnell.", style="List Bullet")
d.add_heading("Überschrift 3", 3)
d.add_paragraph("Sie müssen nichts weiter tun. Wir melden uns bei Ihnen.", style="Quote")

# Tabelle: Kopfzeile Dunkelblau, Zebra Sand
t = d.add_table(rows=3, cols=3); t.autofit = True
for i, row in enumerate(t.rows):
    for j, c in enumerate(row.cells):
        c.text = ["Spalte A", "Spalte B", "Spalte C"][j] if i == 0 else f"Wert {i}.{j+1}"
        fill = C["primary"] if i == 0 else (C["sand"] if i % 2 == 0 else C["white"])
        tcPr = c._tc.get_or_add_tcPr(); sh = OxmlElement("w:shd")
        sh.set(qn("w:val"), "clear"); sh.set(qn("w:fill"), fill); tcPr.append(sh)
        for p in c.paragraphs:
            for run in p.runs:
                run.bold = i == 0
                run.font.color.rgb = rgb("white") if i == 0 else rgb("primary")
d.core_properties.title = "R+V Dokumentvorlage"
d.save(root / "templates/ruv.docx")
print("ok")
